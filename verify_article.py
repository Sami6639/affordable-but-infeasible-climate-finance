"""Check the supplied JCMS Word manuscript against the frozen numerical package.
Usage: python verify_article.py /path/to/manuscript.docx
Requires python-docx in addition to the numerical environment.
Does not upload or copy the manuscript.
"""
from pathlib import Path
import csv, hashlib, json, sys
from docx import Document

ROOT = Path(__file__).resolve().parent
BASE = ROOT / 'revision-2026-10-07/results'

def run(path):
    doc = Document(path)
    checks = []
    def check(label, actual, expected):
        checks.append({'item': label, 'actual': actual, 'expected': expected,
                       'passed': actual == expected})
    def rows(name):
        return list(csv.DictReader((BASE / name).open()))
    def cells(index):
        return [[c.text for c in r.cells] for r in doc.tables[index].rows[1:]]
    def compare_table(label, actual, expected):
        check(label + ' row count', len(actual), len(expected))
        for i, (a, b) in enumerate(zip(actual, expected), 1):
            for j, (x, y) in enumerate(zip(a, b), 1):
                check(f'{label} row {i} column {j}', x, y)
    check('table count including A1', len(doc.tables), 6)
    check('figure count', len(doc.inline_shapes), 4)
    expected = []
    for r in rows('table2_constraint_failures.csv'):
        expected.append([r['case'], r['regime']] +
                        [f'{float(r[k]):.3f}' if r[k] != 'Not required' else r[k]
                         for k in ['lower_margin', 'upper_margin', 'span_margin', 'maximum_obligation']])
    compare_table('Table 2', cells(1), expected)
    expected = []
    for r in rows('table3_advances.csv'):
        expected.append([f"{float(r['cap_ratio']):.2f}",
                         f"{100*float(r['loading']):.0f}%",
                         f"{float(r['fixed_cost_eur']):,.0f}"] +
                        [f'{float(r[k]):,.0f}' if r[k] != 'Infeasible' else r[k]
                         for k in ['exact_advance_eur', 'ceiling_advance_eur']])
    compare_table('Table 3', cells(2), expected)
    expected = []
    for r in rows('table5_dispersion.csv'):
        sd = 100*float(r['relative_sd'])
        sdtext = f'{sd:.1f}' if abs(sd-round(sd,1)) < 1e-9 else f'{sd:.5f}'
        expected.append([sdtext] +
                        ['Pass' if r[k] == 'True' else 'Fail' for k in ['mean_test_90pct','span_test_90pct']] +
                        ['Yes' if r['exact_regime_exists'] == 'True' else 'No',
                         f"{float(r['max_ceiling_advance_percent']):.3f}"])
    compare_table('Table 5', cells(4), expected)
    for i, shape in enumerate(doc.inline_shapes, 1):
        rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        embedded = doc.part.related_parts[rid].blob
        original = (BASE / f'figure{i}.png').read_bytes()
        check(f'Figure {i} SHA256', hashlib.sha256(embedded).hexdigest(),
              hashlib.sha256(original).hexdigest())
    body = '\n'.join(p.text for p in doc.paragraphs)
    for token in ['84.026%', '€473,403.82', '263,119.82', '345,455.97']:
        check('reported result ' + token, token in body, True)
    return {'manuscript_sha256': hashlib.sha256(Path(path).read_bytes()).hexdigest(),
            'checks_passed': sum(x['passed'] for x in checks),
            'checks_total': len(checks), 'checks': checks}

if __name__ == '__main__':
    result = run(sys.argv[1])
    print(json.dumps(result, ensure_ascii=False, indent=2))
    sys.exit(0 if result['checks_passed'] == result['checks_total'] else 1)
